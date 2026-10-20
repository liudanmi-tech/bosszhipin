from pathlib import Path

import yaml

ROOT = Path(__file__).resolve().parent.parent


def load_config(path: str | Path = ROOT / "config.yaml", need_resume: bool = True) -> dict:
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"找不到 {path}，请先 cp config.example.yaml config.yaml 并填写")
    cfg = yaml.safe_load(path.read_text(encoding="utf-8"))

    if not need_resume:
        return cfg
    resume_file = cfg.get("resume_file")
    if resume_file and not cfg.get("resume"):
        resume_path = ROOT / resume_file
        if not resume_path.exists():
            raise FileNotFoundError(f"找不到简历文件 {resume_path}，请创建该文件并粘贴简历纯文本")
        cfg["resume"] = read_resume(resume_path)
    if not cfg.get("resume"):
        raise ValueError("config.yaml 中需要提供 resume 或 resume_file")
    return cfg


def read_resume(path: Path) -> str:
    """读取简历：支持 Word (.docx)、PDF 和纯文本 (.md/.txt)。"""
    if path.suffix.lower() == ".doc":
        raise ValueError("不支持旧版 .doc，请在 Word 里「另存为」.docx 格式")
    if path.suffix.lower() == ".pdf":
        from pypdf import PdfReader

        return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
    if path.suffix.lower() != ".docx":
        return path.read_text(encoding="utf-8")

    from docx import Document

    doc = Document(str(path))
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:  # 很多简历模板用表格排版
        for row in table.rows:
            cells = dict.fromkeys(c.text.strip() for c in row.cells if c.text.strip())
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)
