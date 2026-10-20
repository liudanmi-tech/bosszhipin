import re
from pathlib import Path

from docx import Document
from docx.shared import Pt


def markdown_to_docx(markdown: str, path: Path) -> None:
    """把大模型输出的 Markdown 简历转成 Word：# 标题、## 小节、- 列表、其余为段落。"""
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "PingFang SC"
    style.font.size = Pt(10.5)

    for raw in markdown.splitlines():
        line = re.sub(r"\*\*(.+?)\*\*", r"\1", raw).rstrip()
        if not line.strip() or line.strip() in ("---", "```", "```markdown"):
            continue
        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=0)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
        elif re.match(r"^\s*[-*] ", line):
            doc.add_paragraph(re.sub(r"^\s*[-*] ", "", line), style="List Bullet")
        else:
            doc.add_paragraph(line.strip())

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
